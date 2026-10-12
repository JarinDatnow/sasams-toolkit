"""
SASAMS MCP server: lets Claude answer questions about a copy of the SASAMS
Access database, read-only, without ever seeing a learner's name.

Claude gets two kinds of tools:
- Ready-made tools for the common questions (subject results, term-on-term
  changes, learners at risk, attendance), built on the same queries the Top 10
  and Annexure K generators already use.
- A guarded run_select tool for anything else: SELECT only, academic tables
  only, personal columns blocked (see query_guard.py).

Learners are identified by SASAMS's internal row number (Learner_Info.ID),
which means nothing outside the database. When a teacher needs actual names,
export_learner_list writes them to an Excel file on this PC; the names never
go back to Claude.

Configuration (environment variables, set in the `claude mcp add` command):
  SASAMS_DB_PATH      full path to a COPY of the term's .mdb (never the live file)
  SASAMS_DB_PASSWORD  database password (or put DB_PASSWORD in db_secret.py
                      next to this file, as the other generators do)
  SASAMS_YEAR         default data year (defaults to the current year)
  SASAMS_OUTPUT_DIR   where exported workbooks go (defaults to ./Output)
"""

import datetime
import importlib.util
import os
import sqlite3
import statistics
import sys
from collections import defaultdict
from pathlib import Path

try:
    from mcp.server.mcpserver import MCPServer
except ImportError:  # mcp 1.x
    from mcp.server.fastmcp import FastMCP as MCPServer

from query_guard import MAX_ROWS_HARD_LIMIT, QueryGuard, QueryRejected, load_schema

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
SCHEMA_PATH = REPO / "SASAMS_map.json"

PHASE_BY_GRADE = {
    0: "Foundation", 1: "Foundation", 2: "Foundation", 3: "Foundation",
    4: "Intermediate", 5: "Intermediate", 6: "Intermediate",
    7: "Senior", 8: "Senior", 9: "Senior",
    10: "FET", 11: "FET", 12: "FET",
}
DEFAULT_LEVEL_EDGES = [30.0, 40.0, 50.0, 60.0, 70.0, 80.0]
FLAGGED_LIMIT = 150

INSTRUCTIONS = """\
Read-only access to this school's SASAMS (SA-SAMS) database copy.

Data model: Learner_Info (current learners have Status = 'C'; Class = Classes.ClassId),
Classes (ClassId, Grade, ClassName), Subjects (Id, Name, names carry a grade suffix like
'Mathematics (Gr 10)'), ReportCycles (CycleId, Phase, Term, Datayear; one cycle per
phase per term), ReportMarks (LearnerID, SubjectId, ReportId = CycleId, Mark = the
term's rounded report mark), LearnerCass + SubjectCriteria (individual assessment
tasks), Absentees (one row per learner per day absent). Grades: 0 = Grade R.
Phases: R-3 Foundation, 4-6 Intermediate, 7-9 Senior, 10-12 FET.

Privacy rules: learners are identified only by their internal ID. Names, ID numbers,
contact, medical and comment fields are blocked. Never ask the user to paste names.
When a teacher needs names, call export_learner_list with the IDs: it writes an Excel
file on this PC and only returns the file path.

Prefer the ready-made tools; use run_select for anything they don't cover, and
describe_table / search_schema to find columns first.
"""

mcp = MCPServer("sasams", instructions=INSTRUCTIONS)
guard = QueryGuard(load_schema(SCHEMA_PATH))
_conn = None
_cache = {}


# ---------------------------------------------------------------------------
# Database access (read-only)
# ---------------------------------------------------------------------------
def _password():
    if os.environ.get("SASAMS_DB_PASSWORD"):
        return os.environ["SASAMS_DB_PASSWORD"]
    for folder in (HERE, REPO / "Top 10 automation", REPO / "Annexure K automation"):
        secret = folder / "db_secret.py"
        if secret.exists():
            spec = importlib.util.spec_from_file_location(f"db_secret_{abs(hash(folder))}", secret)
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
            return module.DB_PASSWORD
    raise RuntimeError(
        "No database password found. Set SASAMS_DB_PASSWORD or copy "
        "db_secret.example.py to db_secret.py in the MCP server folder."
    )


def _db_path():
    path = os.environ.get("SASAMS_DB_PATH")
    if not path:
        raise RuntimeError("SASAMS_DB_PATH is not set. Point it at a COPY of the term's .mdb file.")
    if not Path(path).exists():
        raise RuntimeError(f"Database file not found: {path}")
    return path


def _connect():
    test_db = os.environ.get("SASAMS_SQLITE_PATH")
    if test_db:
        return sqlite3.connect(f"file:{test_db}?mode=ro", uri=True, check_same_thread=False)
    import pyodbc

    conn_str = (
        r"DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};"
        rf"DBQ={_db_path()};"
        rf"PWD={_password()};"
        r"ReadOnly=True;"
    )
    return pyodbc.connect(conn_str, readonly=True)


def _query(sql, *params, limit=None):
    global _conn
    for attempt in (1, 2):
        try:
            if _conn is None:
                _conn = _connect()
            cur = _conn.cursor()
            if params:
                cur.execute(sql, params)
            else:
                cur.execute(sql)
            columns = [d[0] for d in cur.description]
            raw = cur.fetchmany(limit) if limit else cur.fetchall()
            return columns, [dict(zip(columns, tuple(r))) for r in raw]
        except Exception:
            _conn = None
            if attempt == 2:
                raise


def _rows(sql, *params):
    return _query(sql, *params)[1]


def _year(year):
    if year:
        return int(year)
    return int(os.environ.get("SASAMS_YEAR") or datetime.date.today().year)


def _grade(grade):
    s = str(grade).strip().upper().replace("GRADE", "").strip()
    g = 0 if s == "R" else int(s)
    if g not in PHASE_BY_GRADE:
        raise ValueError(f"Unknown grade {grade!r}. Use R or 1-12.")
    return g


def _as_date(value):
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    return datetime.date.fromisoformat(str(value)[:10])


def _round(x, places=1):
    return None if x is None else round(float(x), places)


# ---------------------------------------------------------------------------
# Shared lookups
# ---------------------------------------------------------------------------
def _classes():
    if "classes" not in _cache:
        _cache["classes"] = {
            r["ClassId"]: {"grade": r["Grade"], "name": (r["ClassName"] or "").strip()}
            for r in _rows("SELECT ClassId, Grade, ClassName FROM Classes")
        }
    return _cache["classes"]


def _subjects():
    if "subjects" not in _cache:
        _cache["subjects"] = {r["Id"]: (r["Name"] or "").strip() for r in _rows("SELECT Id, Name FROM Subjects")}
    return _cache["subjects"]


def _level_edges():
    if "edges" not in _cache:
        rows = _rows(
            "SELECT MarkFrom FROM Ana2012EvaluationLevels "
            "WHERE EvalVer = 3 AND GradeFrom = 7 AND GradeTo = 12 "
            "AND Code BETWEEN 2 AND 7 ORDER BY Code"
        )
        edges = [float(r["MarkFrom"]) for r in rows]
        _cache["edges"] = edges if len(edges) == 6 else DEFAULT_LEVEL_EDGES
    return _cache["edges"]


def _level(mark):
    for i, edge in enumerate(_level_edges()):
        if mark < edge:
            return i + 1
    return 7


def _cycle_id(grade, term, year):
    phase = PHASE_BY_GRADE[grade]
    rows = _rows(
        "SELECT CycleId FROM ReportCycles WHERE Datayear = ? AND Term = ? AND Phase = ?",
        str(year), int(term), phase,
    )
    if not rows:
        raise ValueError(f"No {phase} report cycle for Term {term} {year}. Use list_report_cycles to see what exists.")
    return rows[0]["CycleId"]


def _learners(grade, class_name=None):
    """Current learners in a grade (optionally one class): {ID: class name}."""
    classes = _classes()
    rows = _rows("SELECT ID, Class FROM Learner_Info WHERE Grade = ? AND Status = 'C'", grade)
    result = {r["ID"]: classes.get(r["Class"], {}).get("name", "?") for r in rows}
    if class_name:
        wanted = class_name.strip().lower()
        result = {lid: c for lid, c in result.items() if c.lower() == wanted}
        if not result:
            names = sorted({c for c in _learners(grade).values()})
            raise ValueError(f"No learners in class {class_name!r} for that grade. Classes: {', '.join(names)}")
    return result


def _marks(cycle_id, learners, subject_contains=None):
    """[(learner_id, subject_id, mark)] for one report cycle, current learners only."""
    subjects = _subjects()
    needle = subject_contains.lower() if subject_contains else None
    out = []
    for r in _rows("SELECT LearnerID, SubjectId, Mark FROM ReportMarks WHERE ReportId = ?", cycle_id):
        if r["LearnerID"] not in learners or r["Mark"] is None:
            continue
        if needle and needle not in subjects.get(r["SubjectId"], "").lower():
            continue
        out.append((r["LearnerID"], r["SubjectId"], float(r["Mark"])))
    return out


def _stats(marks):
    levels = [0] * 7
    for m in marks:
        levels[_level(m) - 1] += 1
    n = len(marks)
    return {
        "learners": n,
        "average": _round(statistics.fmean(marks)) if n else None,
        "median": _round(statistics.median(marks)) if n else None,
        "levels_1_to_7": levels,
        "pct_below_level_2": _round(100 * levels[0] / n) if n else None,
        "pct_level_4_plus": _round(100 * sum(levels[3:]) / n) if n else None,
    }


# ---------------------------------------------------------------------------
# Tools
# ---------------------------------------------------------------------------
@mcp.tool()
def school_overview() -> dict:
    """Grades and classes with the number of current learners in each."""
    classes = _classes()
    counts = defaultdict(int)
    for r in _rows("SELECT Class FROM Learner_Info WHERE Status = 'C'"):
        counts[r["Class"]] += 1
    grades = defaultdict(list)
    for cid, c in sorted(classes.items(), key=lambda kv: (kv[1]["grade"], kv[1]["name"])):
        grades[c["grade"]].append({"class": c["name"], "learners": counts.get(cid, 0)})
    return {
        "level_cutoffs_used": _level_edges(),
        "grades": [
            {"grade": "R" if g == 0 else g, "phase": PHASE_BY_GRADE.get(g), "learners": sum(x["learners"] for x in cl), "classes": cl}
            for g, cl in sorted(grades.items())
        ],
    }


@mcp.tool()
def list_report_cycles(year: int | None = None) -> dict:
    """Report cycles (term x phase) captured for a year; marks only exist for these."""
    y = _year(year)
    rows = _rows("SELECT CycleId, Name, Phase, Term FROM ReportCycles WHERE Datayear = ? ORDER BY Term, Phase", str(y))
    return {"year": y, "cycles": rows}


@mcp.tool()
def subject_results(
    grade: str,
    term: int,
    year: int | None = None,
    class_name: str | None = None,
    subject_contains: str | None = None,
) -> dict:
    """Per-subject results for a grade's term report: learner count, average, median,
    CAPS level distribution (levels 1-7) and the share below level 2 / at level 4+.
    Optional filters: one class (e.g. '10B') and part of a subject name ('math')."""
    g, y = _grade(grade), _year(year)
    learners = _learners(g, class_name)
    by_subject = defaultdict(list)
    for _, sid, mark in _marks(_cycle_id(g, term, y), learners, subject_contains):
        by_subject[sid].append(mark)
    subjects = _subjects()
    results = [{"subject": subjects.get(sid, f"Subject {sid}"), **_stats(m)} for sid, m in by_subject.items()]
    results.sort(key=lambda r: (r["average"] is None, r["average"]))
    return {
        "grade": grade, "term": term, "year": y, "class": class_name,
        "learners_in_selection": len(learners),
        "subjects": results,
    }


@mcp.tool()
def compare_terms(
    grade: str,
    term_from: int,
    term_to: int,
    year: int | None = None,
    class_name: str | None = None,
    subject_contains: str | None = None,
    drop_threshold: float = 10,
) -> dict:
    """How marks changed between two terms of the same year, per subject, plus the
    learners (by ID) whose mark in a subject fell by at least drop_threshold points."""
    g, y = _grade(grade), _year(year)
    learners = _learners(g, class_name)
    before = {(l, s): m for l, s, m in _marks(_cycle_id(g, term_from, y), learners, subject_contains)}
    after = {(l, s): m for l, s, m in _marks(_cycle_id(g, term_to, y), learners, subject_contains)}
    subjects = _subjects()
    per_subject = defaultdict(list)
    flagged = []
    for key in before.keys() & after.keys():
        lid, sid = key
        change = after[key] - before[key]
        per_subject[sid].append(change)
        if change <= -abs(drop_threshold):
            flagged.append({
                "learner_id": lid, "class": learners[lid], "subject": subjects.get(sid, str(sid)),
                f"term_{term_from}": before[key], f"term_{term_to}": after[key], "change": _round(change),
            })
    summary = [
        {
            "subject": subjects.get(sid, str(sid)),
            "learners_compared": len(ch),
            "average_change": _round(statistics.fmean(ch)),
            "dropped_by_threshold_or_more": sum(1 for c in ch if c <= -abs(drop_threshold)),
            "improved_by_threshold_or_more": sum(1 for c in ch if c >= abs(drop_threshold)),
        }
        for sid, ch in per_subject.items()
    ]
    summary.sort(key=lambda r: r["average_change"])
    flagged.sort(key=lambda r: r["change"])
    return {
        "grade": grade, "year": y, "terms": [term_from, term_to], "class": class_name,
        "subjects": summary,
        "flagged_total": len(flagged),
        "flagged": flagged[:FLAGGED_LIMIT],
        "note": "Learners are identified by internal ID. Use export_learner_list to give a teacher their names.",
    }


@mcp.tool()
def at_risk_learners(
    grade: str,
    term: int,
    year: int | None = None,
    class_name: str | None = None,
    below_mark: float = 30,
    min_subjects: int = 2,
) -> dict:
    """Learners (by ID) with at least min_subjects subjects below below_mark in a term."""
    g, y = _grade(grade), _year(year)
    learners = _learners(g, class_name)
    subjects = _subjects()
    low = defaultdict(list)
    for lid, sid, mark in _marks(_cycle_id(g, term, y), learners):
        if mark < below_mark:
            low[lid].append({"subject": subjects.get(sid, str(sid)), "mark": mark})
    risky = [
        {"learner_id": lid, "class": learners[lid], "subjects_below": len(s), "details": sorted(s, key=lambda x: x["mark"])}
        for lid, s in low.items() if len(s) >= min_subjects
    ]
    risky.sort(key=lambda r: -r["subjects_below"])
    by_class = defaultdict(int)
    for r in risky:
        by_class[r["class"]] += 1
    return {
        "grade": grade, "term": term, "year": y, "rule": f"{min_subjects}+ subjects below {below_mark}%",
        "learners_checked": len(learners),
        "at_risk_total": len(risky),
        "at_risk_by_class": dict(sorted(by_class.items())),
        "learners": risky[:FLAGGED_LIMIT],
        "note": "Use export_learner_list with these IDs to give the teacher a named list.",
    }


@mcp.tool()
def learner_marks(learner_id: int, year: int | None = None) -> dict:
    """Every report mark for one learner (by internal ID) across the year's terms."""
    y = _year(year)
    info = _rows("SELECT ID, Grade, Class, Status FROM Learner_Info WHERE ID = ?", learner_id)
    if not info:
        raise ValueError(f"No learner with ID {learner_id}.")
    cycles = {r["CycleId"]: r["Term"] for r in _rows("SELECT CycleId, Term FROM ReportCycles WHERE Datayear = ?", str(y))}
    subjects = _subjects()
    table = defaultdict(dict)
    for r in _rows("SELECT SubjectId, Mark, ReportId FROM ReportMarks WHERE LearnerID = ?", learner_id):
        if r["ReportId"] in cycles and r["Mark"] is not None:
            table[subjects.get(r["SubjectId"], str(r["SubjectId"]))][f"term_{cycles[r['ReportId']]}"] = r["Mark"]
    i = info[0]
    return {
        "learner_id": learner_id, "grade": i["Grade"], "class": _classes().get(i["Class"], {}).get("name"),
        "current": i["Status"] == "C", "year": y,
        "marks": [{"subject": s, **t} for s, t in sorted(table.items())],
    }


@mcp.tool()
def attendance_summary(
    grade: str,
    term: int | None = None,
    year: int | None = None,
    class_name: str | None = None,
    flag_days: int = 5,
) -> dict:
    """Days absent per class for a grade (whole year, or one term), and the learners
    (by ID) absent at least flag_days days."""
    g, y = _grade(grade), _year(year)
    learners = _learners(g, class_name)
    start = end = None
    if term:
        t = _rows("SELECT StartDate, EndDate FROM SchoolTerms WHERE CurrentYear = ? AND Term = ?", str(y), int(term))
        if not t:
            raise ValueError(f"No SchoolTerms row for Term {term} {y}.")
        start, end = _as_date(t[0]["StartDate"]), _as_date(t[0]["EndDate"])
    days = defaultdict(int)
    for r in _rows("SELECT Learnerid, DateAbsent FROM Absentees WHERE DataYear = ?", str(y)):
        lid = r["Learnerid"]
        if lid not in learners:
            continue
        d = _as_date(r["DateAbsent"])
        if start and (d is None or d < start or d > end):
            continue
        days[lid] += 1
    by_class = defaultdict(lambda: {"learners": 0, "days_absent": 0})
    for lid, cls in learners.items():
        by_class[cls]["learners"] += 1
        by_class[cls]["days_absent"] += days.get(lid, 0)
    flagged = sorted(
        ({"learner_id": lid, "class": learners[lid], "days_absent": n} for lid, n in days.items() if n >= flag_days),
        key=lambda r: -r["days_absent"],
    )
    return {
        "grade": grade, "year": y, "term": term,
        "period": [str(start), str(end)] if start else "whole year",
        "by_class": [
            {"class": c, **v, "avg_days_per_learner": _round(v["days_absent"] / v["learners"]) if v["learners"] else None}
            for c, v in sorted(by_class.items())
        ],
        "flagged_total": len(flagged),
        "flagged": flagged[:FLAGGED_LIMIT],
    }


@mcp.tool()
def export_learner_list(learner_ids: list[int], title: str) -> dict:
    """Write an Excel list (learner number, surname, name, grade, class) for these
    learner IDs to this PC, for a teacher. Names stay on this PC: only the file path
    and a count come back."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    ids = sorted({int(i) for i in learner_ids})
    if not ids:
        raise ValueError("No learner IDs given.")
    if len(ids) > 2000:
        raise ValueError("Too many learners for one list (max 2000).")
    classes = _classes()
    rows = []
    for lid in ids:
        found = _rows("SELECT ID, AccessionNo, SName, FName, NickName, Grade, Class FROM Learner_Info WHERE ID = ?", lid)
        if found:
            r = found[0]
            rows.append([
                str(r["AccessionNo"] or "").strip(),
                (r["SName"] or "").strip().upper(),
                (r["NickName"] or r["FName"] or "").strip().upper(),
                "R" if r["Grade"] == 0 else r["Grade"],
                classes.get(r["Class"], {}).get("name", ""),
            ])
    rows.sort(key=lambda r: (str(r[3]), r[4], r[1], r[2]))

    out_dir = Path(os.environ.get("SASAMS_OUTPUT_DIR") or HERE / "Output")
    out_dir.mkdir(parents=True, exist_ok=True)
    safe = "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in title).strip() or "Learner list"
    path = out_dir / f"{safe} {datetime.datetime.now():%Y-%m-%d %H%M}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Learners"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {datetime.date.today():%d %B %Y} from SASAMS. Contains personal information: handle under POPIA."])
    ws.append([])
    ws.append(["Learner no.", "Surname", "Name", "Grade", "Class"])
    for cell in ws[4]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for col, width in zip("ABCDE", (14, 24, 20, 8, 10)):
        ws.column_dimensions[col].width = width
    wb.save(path)
    return {"file": str(path), "learners_written": len(rows), "not_found": len(ids) - len(rows)}


def _load_generator(folder, filename, name):
    spec = importlib.util.spec_from_file_location(name, REPO / folder / filename)
    module = importlib.util.module_from_spec(spec)
    sys.path.insert(0, str(REPO / folder))
    try:
        spec.loader.exec_module(module)
    finally:
        sys.path.pop(0)
    return module


@mcp.tool()
def generate_top10(term: int, year: int | None = None) -> dict:
    """Run the existing Top 10 generator for a term. The workbook (with names) is
    saved on this PC; only its path comes back."""
    module = _load_generator("Top 10 automation", "generate_top10.py", "generate_top10")
    path = module.generate(db_path=_db_path(), db_password=_password(), term=int(term), year=_year(year))
    return {"file": str(path)}


@mcp.tool()
def generate_annexure_k(term: int, year: int | None = None) -> dict:
    """Run the existing Annexure K generator for a term (8-12 and R-7 groupings, using
    the templates configured in that script). Only the saved file paths come back."""
    module = _load_generator("Annexure K automation", "generate_annexure_k.py", "generate_annexure_k")
    paths = module.generate(db_path=_db_path(), db_password=_password(), term=int(term), data_year=_year(year))
    return {"files": [str(p) for p in paths]}


@mcp.tool()
def search_schema(keyword: str) -> dict:
    """Find tables and columns whose names contain a keyword. Shows which tables
    run_select can use and hides personal columns."""
    k = keyword.lower()
    hits = []
    for table, cols in guard.schema.items():
        queryable = guard.table_is_queryable(table)
        visible = [c for c in cols if not (queryable and guard.column_is_hidden(table, c))]
        matching = [c for c in visible if k in c.lower()]
        if k in table.lower() or matching:
            hits.append({"table": table, "queryable": queryable, "matching_columns": matching[:30]})
    hits.sort(key=lambda h: (not h["queryable"], h["table"]))
    return {"keyword": keyword, "results": hits[:60]}


@mcp.tool()
def describe_table(table: str) -> dict:
    """Columns of one table, with personal columns marked as hidden."""
    match = next((t for t in guard.schema if t.lower() == table.lower()), None)
    if not match:
        raise ValueError(f"No table named {table!r}. Try search_schema.")
    queryable = guard.table_is_queryable(match)
    cols = guard.schema[match]
    return {
        "table": match,
        "queryable": queryable,
        "columns": [c for c in cols if not guard.column_is_hidden(match, c)] if queryable else cols,
        "hidden_columns": sum(1 for c in cols if guard.column_is_hidden(match, c)) if queryable else None,
    }


@mcp.tool()
def run_select(sql: str, max_rows: int = 200) -> dict:
    """Run one read-only SELECT on the academic tables (Access SQL: TOP n instead of
    LIMIT, wildcards * in LIKE, IIF instead of CASE). Personal columns and SELECT *
    are rejected with a reason; fix the query and try again."""
    try:
        clean = guard.check_sql(sql)
    except QueryRejected as e:
        return {"error": str(e)}
    limit = max(1, min(int(max_rows), MAX_ROWS_HARD_LIMIT))
    try:
        columns, rows = _query(clean, limit=limit + 1)
    except Exception as e:  # show the database's own message so Claude can fix the SQL
        return {"error": f"Database error: {e}"}
    truncated = len(rows) > limit
    rows, hidden = guard.redact_rows(columns, rows[:limit])
    for row in rows:
        for key, value in row.items():
            if isinstance(value, (datetime.date, datetime.datetime)):
                row[key] = value.isoformat()
            elif not isinstance(value, (int, float, str, bool, type(None))):
                row[key] = str(value)
    return {"columns": columns, "rows": rows, "row_count": len(rows), "truncated": truncated, "hidden_columns": hidden}


if __name__ == "__main__":
    mcp.run()
