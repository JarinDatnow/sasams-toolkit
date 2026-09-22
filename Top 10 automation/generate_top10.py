"""
Generates the school's "Top 10" achiever workbook straight from the Access
database, replacing the manual process of exporting 28 reports from d6 by hand.

Requires: Python 3 (64-bit), pyodbc, openpyxl, and the 64-bit
"Microsoft Access Database Engine" ODBC driver (already installed on this
machine - if you ever run this on a different PC and get a driver error,
install the free "Microsoft Access Database Engine 2016 Redistributable",
64-bit version, from Microsoft).

HOW TO RUN NEXT TERM
1. Copy the new term's Access database (.mdb) into this folder, or update
   DB_PATH below to point at it. Always use a COPY, never the live file.
2. Update TERM and YEAR below.
3. Run:  python generate_top10.py
4. The finished workbook appears in the OUTPUT_FOLDER, and a summary is
   printed to the screen - check it for any ties or warnings.
"""

import copy
import datetime
import os
from decimal import ROUND_HALF_UP, Decimal

import pyodbc
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

try:
    # The real password lives in db_secret.py, which is gitignored and never
    # committed. Copy db_secret.example.py to db_secret.py and fill it in.
    from db_secret import DB_PASSWORD
except ImportError:
    raise RuntimeError(
        "db_secret.py not found. Copy db_secret.example.py to db_secret.py "
        "in this folder and put the real database password in it."
    )

# ---------------------------------------------------------------------------
# SETTINGS - change these each term
# ---------------------------------------------------------------------------
DB_PATH = (
    r"C:\Users\user\Desktop\Jarin Werk\2026\Claude playground\Top 10 automation"
    r"\700400012 Balmoral College 22 Sept\700400012 Balmoral College.mdb"
)
TEMPLATE_PATH = (
    r"C:\Users\user\Desktop\Jarin Werk\2026\Claude playground\Top 10 automation"
    r"\TEMPLATE  top 10.xlsx"
)
TERM = 3
YEAR = 2026
OUTPUT_FOLDER = (
    r"C:\Users\user\Desktop\Jarin Werk\2026\Claude playground\Top 10 automation\Output"
)

# ---------------------------------------------------------------------------
# Fixed layout - matches TEMPLATE__top_10.xlsx, should not normally change
# ---------------------------------------------------------------------------
TEMPLATE_SHEET_NAME = "Worksheet"
FIRST_DATA_ROW = 13
LAST_TEMPLATE_DATA_ROW = 22  # template ships 10 pre-styled rows: 13-22
DATA_ROW_HEIGHT = 28.5
COLUMNS = {"nr": "A", "learner_no": "B", "surname": "C", "name": "D", "pct": "E"}
TOP_N = 10

# Grade -> report-cycle phase, per the CAPS phase structure this school uses
PHASE_BY_GRADE = {
    0: "Foundation", 1: "Foundation", 2: "Foundation", 3: "Foundation",
    4: "Intermediate", 5: "Intermediate", 6: "Intermediate",
    7: "Senior", 8: "Senior", 9: "Senior",
    10: "FET", 11: "FET", 12: "FET",
}

# Long-name shrink threshold: names longer than this (in the SURNAME or
# PREFERRED NAME column) get shrink-to-fit turned on so they don't clip
# when printed.
SHRINK_NAME_LEN = 11

# Sanity-check window for task dates, see check_task_dates() below.
TASK_DATE_WARNING_BUFFER_DAYS = 14


def connect():
    conn_str = (
        r"DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};"
        rf"DBQ={DB_PATH};"
        rf"PWD={DB_PASSWORD};"
        r"ReadOnly=True;"
    )
    return pyodbc.connect(conn_str, readonly=True)


def round_half_up(value):
    return int(Decimal(repr(value)).quantize(0, rounding=ROUND_HALF_UP))


def get_report_id_map(cursor):
    """phase name -> ReportId (CycleId) for the configured TERM/YEAR."""
    cursor.execute(
        "SELECT Phase, CycleId FROM ReportCycles WHERE Datayear = ? AND Term = ?",
        str(YEAR), TERM,
    )
    mapping = {row.Phase: row.CycleId for row in cursor.fetchall()}
    missing = set(PHASE_BY_GRADE.values()) - set(mapping)
    if missing:
        raise RuntimeError(
            f"No ReportCycles row found for Term {TERM} {YEAR}, phase(s): {missing}. "
            "Check TERM/YEAR are correct and that report cycle exists in the database."
        )
    return mapping


def get_classes(cursor):
    """(grade, classid, classname) for grades R(0)-6, ordered for output."""
    cursor.execute(
        "SELECT ClassId, Grade, ClassName FROM Classes WHERE Grade BETWEEN 0 AND 6 "
        "ORDER BY Grade, ClassName"
    )
    return [(row.Grade, row.ClassId, row.ClassName) for row in cursor.fetchall()]


def get_term_dates(cursor):
    """(StartDate, EndDate) for the configured TERM/YEAR, from SchoolTerms.
    Used only for the check_task_dates() sanity check below - the actual
    mark calculation matches tasks to the term via SubjectCriteria.SubHeading
    (see unrounded_subject_pct), not by date."""
    cursor.execute(
        "SELECT StartDate, EndDate FROM SchoolTerms WHERE CurrentYear = ? AND Term = ?",
        str(YEAR), TERM,
    )
    row = cursor.fetchone()
    if row is None:
        raise RuntimeError(
            f"No SchoolTerms row found for Term {TERM} {YEAR}. "
            "Check TERM/YEAR are correct."
        )
    return row.StartDate, row.EndDate


def check_task_dates(cursor, term_start, term_end):
    """
    Sanity check, not part of the mark calculation itself: list every
    assessment task tagged as this term (SubjectCriteria.SubHeading =
    'TermN') whose own DateAdded falls more than TASK_DATE_WARNING_BUFFER_DAYS
    outside the term's official [StartDate, EndDate] window (from
    SchoolTerms). A task like that isn't excluded from the calculation -
    SubHeading is trusted as the authoritative term link - but a mismatch
    this large usually means either the task was captured very late/early,
    or (as with a "SBA Year Mark" rollup row) it's a summary task rather
    than a real per-term assessment, which is worth a human look.
    """
    buffer = datetime.timedelta(days=TASK_DATE_WARNING_BUFFER_DAYS)
    lower, upper = term_start - buffer, term_end + buffer
    cursor.execute(
        """
        SELECT sc.Subjectid, s.Name, sc.CriterionID, sc.Description, sc.DateAdded, sc.Weighting
        FROM SubjectCriteria sc LEFT JOIN Subjects s ON sc.Subjectid = s.Id
        WHERE sc.DataYear = ? AND sc.SubHeading = ?
          AND (sc.DateAdded < ? OR sc.DateAdded > ?)
        ORDER BY sc.DateAdded
        """,
        str(YEAR), f"Term{TERM}", lower, upper,
    )
    return cursor.fetchall()


def unrounded_subject_pct(cursor, learner_id, subject_id):
    """
    d6's own PERCENTAGE column is built from the *unrounded* weighted
    average of each subject's assessment tasks for the term, not from the
    already-rounded whole-number ReportMarks.Mark. Reproduce that: each row
    in LearnerCass is one task (test, assignment, exam...) with a Mark out
    of Criterionscore, and SubjectCriteria.Weighting says how much that task
    counts toward the subject total. Tasks are matched to this term via
    SubjectCriteria.SubHeading ('Term1'/'Term2'/'Term3'/'Term4') - an
    explicit term label, not a date guess.
    Returns None if no task breakdown exists for this subject this term.
    """
    cursor.execute(
        """
        SELECT lc.Mark, lc.Criterionscore, sc.Weighting
        FROM LearnerCass lc
        INNER JOIN SubjectCriteria sc
            ON lc.CriterionId = sc.CriterionID
           AND lc.Subjectid = sc.Subjectid
           AND lc.Datayear = sc.DataYear
        WHERE lc.Learnerid = ? AND lc.Subjectid = ? AND lc.Datayear = ?
          AND sc.SubHeading = ?
        """,
        learner_id, subject_id, str(YEAR), f"Term{TERM}",
    )
    rows = cursor.fetchall()
    if not rows:
        return None
    total_weight = sum(r.Weighting for r in rows)
    if not total_weight:
        return None
    weighted = sum(
        (r.Mark / r.Criterionscore) * r.Weighting for r in rows if r.Criterionscore
    )
    return weighted / total_weight * 100


def fetch_learners(cursor, grade, class_id, report_id):
    """Return sorted list of learner dicts for one class/grade report."""
    if class_id is None:
        cursor.execute(
            "SELECT ID, AccessionNo, SName, NickName, FName FROM Learner_Info "
            "WHERE Grade = ? AND Status = 'C'",
            grade,
        )
    else:
        cursor.execute(
            "SELECT ID, AccessionNo, SName, NickName, FName FROM Learner_Info "
            "WHERE Grade = ? AND Class = ? AND Status = 'C'",
            grade, class_id,
        )
    learners = cursor.fetchall()

    results = []
    for learner in learners:
        cursor.execute(
            "SELECT SubjectId, Mark FROM ReportMarks WHERE LearnerID = ? AND ReportId = ?",
            learner.ID, report_id,
        )
        subject_marks = cursor.fetchall()
        if not subject_marks:
            continue  # no marks captured for this learner this term - skip

        subject_pcts = []
        for sm in subject_marks:
            pct = unrounded_subject_pct(cursor, learner.ID, sm.SubjectId)
            subject_pcts.append(pct if pct is not None else sm.Mark)

        avg = sum(subject_pcts) / len(subject_pcts)
        preferred_name = (learner.NickName or learner.FName or "").strip().upper()
        results.append({
            "learner_no": str(learner.AccessionNo).strip(),
            "surname": (learner.SName or "").strip().upper(),
            "name": preferred_name,
            "avg": avg,
            "pct": round_half_up(avg),
        })

    results.sort(key=lambda r: (-r["avg"], r["surname"], r["name"]))
    return results


def top_10_with_ties(learners):
    if len(learners) <= TOP_N:
        return learners, False
    cutoff_pct = learners[TOP_N - 1]["pct"]
    extended = list(learners[:TOP_N])
    had_tie = False
    for learner in learners[TOP_N:]:
        if learner["pct"] == cutoff_pct:
            extended.append(learner)
            had_tie = True
        else:
            break
    return extended, had_tie


def style_row_like(ws, source_row, target_row):
    for col in COLUMNS.values():
        src_cell = ws[f"{col}{source_row}"]
        dst_cell = ws[f"{col}{target_row}"]
        dst_cell.font = copy.copy(src_cell.font)
        dst_cell.border = copy.copy(src_cell.border)
        dst_cell.fill = copy.copy(src_cell.fill)
        dst_cell.alignment = copy.copy(src_cell.alignment)
        dst_cell.number_format = src_cell.number_format
    ws.row_dimensions[target_row].height = DATA_ROW_HEIGHT


def fill_sheet(ws, title_c4, term_c6, learners):
    ws["C4"] = title_c4
    ws["C6"] = term_c6

    last_row = FIRST_DATA_ROW + len(learners) - 1
    # extend styled rows beyond the template's 10 pre-built rows if there are ties
    for row in range(LAST_TEMPLATE_DATA_ROW + 1, last_row + 1):
        style_row_like(ws, LAST_TEMPLATE_DATA_ROW, row)
    # clear any unused template rows (fewer than 10 learners)
    for row in range(FIRST_DATA_ROW + len(learners), LAST_TEMPLATE_DATA_ROW + 1):
        for col in COLUMNS.values():
            ws[f"{col}{row}"] = None

    for i, learner in enumerate(learners):
        row = FIRST_DATA_ROW + i
        ws[f"{COLUMNS['nr']}{row}"] = i + 1
        ws[f"{COLUMNS['learner_no']}{row}"] = learner["learner_no"]
        ws[f"{COLUMNS['surname']}{row}"] = learner["surname"]
        ws[f"{COLUMNS['name']}{row}"] = learner["name"]
        ws[f"{COLUMNS['pct']}{row}"] = learner["pct"]

        for col_key in ("surname", "name"):
            cell = ws[f"{COLUMNS[col_key]}{row}"]
            text = str(cell.value or "")
            if len(text) > SHRINK_NAME_LEN:
                a = copy.copy(cell.alignment)
                a.shrink_to_fit = True
                cell.alignment = a

    # make sure the sheet prints on exactly one page regardless of row count
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{get_column_letter(5)}{max(last_row, LAST_TEMPLATE_DATA_ROW)}"


def generate(db_path=None, db_password=None, template_path=None, term=None,
             year=None, output_folder=None):
    """Run a full Top 10 generation. Any argument left as None falls back to
    the module-level SETTINGS above. Returns the path of the saved workbook.
    Used both by main() (CLI run with the settings at the top of this file)
    and by the web UI (webapp/server.py), which passes its own values in."""
    global DB_PATH, DB_PASSWORD, TEMPLATE_PATH, TERM, YEAR, OUTPUT_FOLDER
    if db_path is not None:
        DB_PATH = db_path
    if db_password is not None:
        DB_PASSWORD = db_password
    if template_path is not None:
        TEMPLATE_PATH = template_path
    if term is not None:
        TERM = term
    if year is not None:
        YEAR = year
    if output_folder is not None:
        OUTPUT_FOLDER = output_folder

    os.makedirs(OUTPUT_FOLDER, exist_ok=True)
    cnxn = connect()
    cursor = cnxn.cursor()

    report_ids = get_report_id_map(cursor)
    classes = get_classes(cursor)
    term_start, term_end = get_term_dates(cursor)
    date_warnings = check_task_dates(cursor, term_start, term_end)

    template_wb = load_workbook(TEMPLATE_PATH)
    template_ws = template_wb[TEMPLATE_SHEET_NAME]

    summary = []  # (sheet name, learner count, tie count)

    # --- Grade R to 6: one sheet per class ---
    for grade, class_id, class_name in classes:
        report_id = report_ids[PHASE_BY_GRADE[grade]]
        learners = fetch_learners(cursor, grade, class_id, report_id)
        final_list, had_tie = top_10_with_ties(learners)

        ws = template_wb.copy_worksheet(template_ws)
        ws.title = f"Grade {class_name}"
        fill_sheet(ws, f"GRADE  {class_name}", f"TERM {TERM}", final_list)

        extra = len(final_list) - min(TOP_N, len(final_list))
        summary.append((ws.title, len(final_list), extra))

    # --- Grade 7 to 12: one sheet per grade, classes combined ---
    for grade in range(7, 13):
        report_id = report_ids[PHASE_BY_GRADE[grade]]
        learners = fetch_learners(cursor, grade, None, report_id)
        final_list, had_tie = top_10_with_ties(learners)

        ws = template_wb.copy_worksheet(template_ws)
        ws.title = f"Grade {grade}"
        fill_sheet(ws, f"GRADE  {grade}", f"TERM {TERM}", final_list)

        extra = len(final_list) - min(TOP_N, len(final_list))
        summary.append((ws.title, len(final_list), extra))

    # remove the template's own sheets
    del template_wb[TEMPLATE_SHEET_NAME]
    if "Worksheet 1" in template_wb.sheetnames:
        del template_wb["Worksheet 1"]

    cnxn.close()

    out_name = f"Top 10 - Term {TERM} {YEAR}.xlsx"
    out_path = os.path.join(OUTPUT_FOLDER, out_name)
    template_wb.save(out_path)

    print(f"\nSaved: {out_path}")
    print(f"Generated {len(summary)} sheets:\n")
    print(f"{'Sheet':<14}{'Learners':>10}{'Extra tie rows':>16}")
    for name, count, extra in summary:
        tie_note = f"+{extra}" if extra else ""
        print(f"{name:<14}{count:>10}{tie_note:>16}")

    ties = [(n, e) for n, _, e in summary if e]
    if ties:
        print("\nSheets with ties beyond rank 10:")
        for n, e in ties:
            print(f"  {n}: {e} extra learner(s) tied with 10th place")
    else:
        print("\nNo ties beyond rank 10 on any sheet.")

    if date_warnings:
        print(
            f"\nWARNING: {len(date_warnings)} task(s) are tagged Term {TERM} "
            f"but their own date is more than {TASK_DATE_WARNING_BUFFER_DAYS} days "
            f"outside the term's official {term_start:%Y-%m-%d} to {term_end:%Y-%m-%d} "
            "window. They ARE still included in the calculation (SubHeading is "
            "trusted as authoritative) - this is just a heads-up in case one is "
            "mistagged or is a year-end rollup task rather than a real term task."
        )
        print(f"{'Subject':<40}{'Task':<35}{'Date':<12}{'Weight':>7}")
        for w in date_warnings:
            print(f"{(w.Name or ''):<40}{w.Description:<35}{w.DateAdded:%Y-%m-%d}  {w.Weighting:>6.1f}")
    else:
        print(f"\nNo Term {TERM} tasks dated outside the term window - nothing to flag.")

    print(f"\nRun finished: {datetime.datetime.now():%Y-%m-%d %H:%M:%S}")

    return out_path


def main():
    generate()


if __name__ == "__main__":
    main()
